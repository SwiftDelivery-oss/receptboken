import contextlib
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from common import assert_no_money
from convert_excel import convert, FormulaReader
from build import validate
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula


class Links(HTMLParser):
    def __init__(self): super().__init__(); self.links=[]; self.robots=False; self.lang=False
    def handle_starttag(self, tag, attrs):
        attrs=dict(attrs)
        if tag=='html': self.lang=attrs.get('lang')=='sv'
        if tag=='meta' and attrs.get('name')=='robots': self.robots='noindex' in attrs.get('content','')
        for key in ('href','src'):
            if attrs.get(key): self.links.append(attrs[key])


class Recipes(unittest.TestCase):
    def test_economic_fields_are_rejected(self):
        for key in ['pris_per_kg','costPerRow','price','kostnad','kr_kg']:
            with self.assertRaises(ValueError): assert_no_money({key: 123})

    def test_data_integrity(self):
        recipes=[json.loads(x.read_text(encoding='utf-8')) for x in (ROOT/'recept').glob('*.json')]
        self.assertGreater(len(recipes),0)
        for r in recipes:
            assert_no_money(r)
            validate(r)
        for slug in ['surdeg-vete','surdeg-rag','skollning','skollning-ragmjol-fint']:
            r=json.loads((ROOT/'recept'/f'{slug}.json').read_text(encoding='utf-8'))
            self.assertAlmostEqual(sum(i['mängd'] for g in r['ingredienser'] for i in g['ingredienser']),1)

    def test_converter_column_detection_array_formula_and_money(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); w=openpyxl.Workbook(); order=w.active; order.title='antal att baka'; order['B5']='Testbröd'
            s=w.create_sheet('Testbröd'); s['B3']='Testbröd'; s['B4']=10; s['C4']='st'
            s['B6']=.025; s['C6']='kg'; s['D6']='Mjöl'
            s['F6']=ArrayFormula(ref='F6',text='=SUM(B6*$B$4)'); s['G6']='kg'
            s['H5']='kr kg'; s['H6']=199; s['I6']='=F6*H6'; s['B8']='Blanda allt'
            s['B9']='Pris 99 kr'; w.create_sheet('Sheet1'); w.save(root/'source.xlsx')
            with contextlib.redirect_stdout(io.StringIO()): convert(root/'source.xlsx',root/'recipes',root/'rapport.md')
            r=json.loads((root/'recipes/testbrod.json').read_text(encoding='utf-8'))
            self.assertEqual(r['ingredienser'][0]['ingredienser'][0]['mängd'],.025)
            self.assertEqual(r['anteckningar'],['Blanda allt'])
            assert_no_money(r)
            with self.assertRaises(FileExistsError):
                with contextlib.redirect_stdout(io.StringIO()): convert(root/'source.xlsx',root/'recipes',root/'rapport2.md')

    def test_formula_reader_rejects_code_and_cycles(self):
        w=openpyxl.Workbook(); s=w.active; s.title='Test'
        s['A1']=2; s['A2']='=SUM(A1*3)'; s['A3']='=A3'; s['A4']='=__import__("os")'
        reader=FormulaReader(w,w)
        self.assertEqual(reader.read('Test','A2'),6)
        with self.assertRaises(ValueError): reader.read('Test','A3')
        with self.assertRaises(ValueError): reader.read('Test','A4')

    def test_public_build_links_noindex_and_money(self):
        dest=ROOT/'dist'; self.assertTrue((dest/'index.html').exists(),'Kör bygget först')
        base=json.loads((dest/'manifest.webmanifest').read_text(encoding='utf-8'))['scope']
        for p in dest.rglob('*'):
            if not p.is_file() or p.suffix not in {'.html','.json','.js','.css','.txt','.webmanifest','.svg'}: continue
            text=p.read_text(encoding='utf-8')
            self.assertIsNone(re.search(r'\b(?:kr|SEK|kronor|kostnad\w*|pris(?:er|fält)?)\b',text,re.I),p)
            if p.suffix=='.html':
                parser=Links(); parser.feed(text)
                self.assertTrue(parser.lang,p); self.assertTrue(parser.robots,p)
                for link in parser.links:
                    if link.startswith('#'): continue
                    self.assertTrue(link.startswith(base),(p,link,base))
                    path=link.split('?')[0].split('#')[0][len(base):]; target=dest/path
                    if link.split('?')[0].split('#')[0].endswith('/'): target/='index.html'
                    self.assertTrue(target.is_file(),(p,link))
        self.assertEqual(len(list((dest/'recept').glob('*/index.html'))),len(list((ROOT/'recept').glob('*.json'))))


if __name__=='__main__': unittest.main()
